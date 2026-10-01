import pathlib
from typing import ClassVar

import polars as pl
from openpyxl import load_workbook

from sharkadm.data.data_holder import PolarsDataHolder
from sharkadm.data.data_source.txt_file import CsvRowFormatPolarsDataFile
from sharkadm.data.data_source.xlsx_file import XlsxFormatPolarsDataFile
from sharkadm.sharkadm_logger import adm_logger

INTERNAL_COLUMN_NAMES = {
    "datetime": "sample_iso_datetime",
    # "LOCAL_CDI_ID": "",
    # "Cruise": "",
    # "Type": "",
    # "dataset_ID": "",
    # "Abstract": "",
    # "documentation_URL":"",
    "start_date": "visit_date",
    "start_time": "sample_time",
    "end_date": "sample_enddate",
    "end_time": "sample_endtime",
    # "Sea_area_code"	: "",
    "longitude": "visit_reported_longitude",
    "latitude": "visit_reported_latitude",
    # "Horizontal_datum"	: "",
    # "Horizontal_datum_code"	: "",
    # "P02 codes"	: "",
    # "platform_type"	: "",
    "station_name": "reported_station_name",
    # "station_short_name"	: "",
    # "station_start_date"	: "",
    # "Measuring_area_type"	: "",
    # "Time_resolution"	: "",
    # "Time_resolution_unit"	: "",
    # "Instrument"	: "",
    # "originator"	: "",
    # "custodian"	: "",
    # "distributor"	: "",
    # "CDI_partner"	: "",
    # "EDMERP"	: "",
    # "revision_date": ""
}


class NewDataHolder(PolarsDataHolder):
    _data_structure = "column"
    _data_set_name = ""

    _indexing_column = "sample_iso_datetime"

    # TODO verify this mapping
    _sdn_qc_mapping: ClassVar[dict[str, str]] = {"0": "1", "1": "2"}

    def __init__(
        self,
        data_root_directory: str | pathlib.Path | None = None,
        data_file_name: str | pathlib.Path | None = None,
        metadata_file_name: str | pathlib.Path | None = None,
        **kwargs,
    ):
        super().__init__(**kwargs)
        self._data_root_directory = pathlib.Path(data_root_directory)
        if not self._data_root_directory.is_dir():
            raise NotADirectoryError(self._data_root_directory)

        self._data_file_name = data_file_name
        self._metadata_file_name = metadata_file_name

        self._data: pl.DataFrame = pl.DataFrame()
        self._data_columns: list[str] = []
        self._dataset_name = self._data_root_directory.name
        self._load_data()

    @staticmethod
    def get_data_holder_description() -> str:
        return """Holds data from [your description here]"""

    @property
    def dataset_name(self) -> str:
        return self._dataset_name or "Unknown"

    @property
    def data_file_path(self) -> pathlib.Path:
        return self._data_root_directory / self._data_file_name

    @property
    def metadata_file_path(self) -> pathlib.Path:
        return self._data_root_directory / self._metadata_file_name

    @property
    def indexing_column(self) -> str:
        """The column that indexes the series."""
        return self._indexing_column

    @property
    def data_columns(self) -> tuple[str, ...]:
        """Columns with data."""
        return tuple(self._data_columns)

    def _update_data_columns(self) -> None:
        excluded_prefixes = ("qc_", "re_")
        self._data_columns = [
            c
            for c in self._data.columns
            if c != self.indexing_column
            and not c.startswith(excluded_prefixes)
            and c != "source"
        ]

    def _load_data(self) -> None:
        data_source = CsvRowFormatPolarsDataFile(
            path=self.data_file_path,
            data_type=self.data_type,
            nodc_conf=self.config,
            delimiter=";",
        )
        data_source.map_header(INTERNAL_COLUMN_NAMES)
        self._set_data_source(data_source)
        self._update_data_columns()
        self._add_sdn_qc_flags()
        self._load_metadata()

    def _load_metadata(self) -> None:
        """Load metadata from either CSV or XLSX file."""
        metadata_path = self.metadata_file_path

        if not metadata_path.exists():
            adm_logger.log_workflow(
                f"Metadata file not found: {metadata_path}", level=adm_logger.WARNING
            )
            return

        if metadata_path.suffix.lower() == ".xlsx":
            workbook = load_workbook(metadata_path, read_only=True, data_only=True)
            sheet_name = workbook.sheetnames[0]
            workbook.close()
            metadata_source = XlsxFormatPolarsDataFile(
                path=metadata_path,
                sheet_name=sheet_name,
                data_type=self.data_type,
                nodc_conf=self.config,
            )
        elif metadata_path.suffix.lower() in [".csv", ".txt"]:
            metadata_source = CsvRowFormatPolarsDataFile(
                path=metadata_path, data_type=self.data_type, nodc_conf=self.config
            )
        else:
            adm_logger.log_workflow(
                f"Unsupported metadata file format: {metadata_path.suffix}",
                level=adm_logger.ERROR,
            )
            return

        metadata_source.map_header(INTERNAL_COLUMN_NAMES)
        metadata_df = self._get_data_from_data_source(metadata_source)

        if len(metadata_df) >= 1:
            metadata_dict = metadata_df[0, :].to_dicts()[0]

            for col_name, value in metadata_dict.items():
                if col_name not in self._data.columns:
                    self._data = self._data.with_columns(pl.lit(value).alias(col_name))

    def _add_sdn_qc_flags(self) -> None:
        """Add SeaDataNet qualifier columns for each measurement column."""
        expressions = []

        for parameter in self._data_columns:
            qc_column = f"qc_{parameter}"
            sdn_qc_column = f"qc_sdn_{parameter}"

            if qc_column not in self._data.columns:
                continue

            expressions.append(
                pl.col(qc_column)
                .cast(pl.String)
                .fill_null("")
                .replace_strict(
                    self._sdn_qc_mapping,
                    default="9",
                )
                .alias(sdn_qc_column)
            )

        if expressions:
            self._data = self._data.with_columns(expressions)
